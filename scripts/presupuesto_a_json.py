# -*- coding: utf-8 -*-
"""Convierte "PRESUPUESTO DE HORAS 2026.xlsx" en data/presupuesto-2026.json.

Uso:  python scripts/presupuesto_a_json.py "ruta/al/PRESUPUESTO DE HORAS 2026.xlsx"

La hoja tiene: fila 1 = meses (celdas combinadas), fila 2 = cargabilidad
(horas disponibles por semana), fila 3 = etiquetas de semana, filas 4+ =
cliente / funcionario / cargo y las horas por semana. La fase de cada celda
viene dada por el color de relleno, según la leyenda al pie de la hoja.
"""
import sys, json, datetime, unicodedata
import openpyxl

PHASE_BY_FILL = {
    'theme8': 'previas',        # Actividades Previas
    'theme5': 'vacaciones',     # Semana de Vacaciones
    'theme6': 'planeacion',     # Planeación
    'FFF5C6F5': 'ejecucion',    # Ejecución
    'FFAEF5F0': 'cierre',       # Cierre
    'theme9': 'informe',        # Informe Mensual
    'FFFFFF00': 'compensatorio',
    'theme7': 'especiales',     # Proyectos Especiales
}
FIRST_WEEK = datetime.date(2026, 6, 1)   # lunes 1 de junio de 2026
FIRST_WEEK_COL = 4                        # columna D


def fill_key(cell):
    f = cell.fill
    if not f or not f.fill_type or not f.fgColor:
        return None
    c = f.fgColor
    if c.type == 'rgb':
        return c.rgb
    if c.type == 'theme':
        return 'theme%s' % c.theme
    return None


def clean(s):
    return ' '.join(str(s).replace('\u00a0', ' ').split()).strip()


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]

    # Meses: celdas combinadas de la fila 1
    month_of_col = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_row == 1:
            name = clean(ws.cell(1, rng.min_col).value)
            for c in range(rng.min_col, rng.max_col + 1):
                month_of_col[c] = name

    weeks = []
    col = FIRST_WEEK_COL
    i = 0
    while ws.cell(3, col).value:
        weeks.append({
            'index': i,
            'start': (FIRST_WEEK + datetime.timedelta(days=7 * i)).isoformat(),
            'label': clean(ws.cell(3, col).value),
            'month': month_of_col.get(col, ''),
            'capacity': ws.cell(2, col).value or 0,
        })
        col += 1
        i += 1

    people, clients, rows = {}, [], []
    row = 4
    while True:
        client = ws.cell(row, 1).value
        person = ws.cell(row, 2).value
        role = ws.cell(row, 3).value
        if not client and not person:
            break
        client, person, role = clean(client), clean(person), clean(role).lower()
        if client not in clients:
            clients.append(client)
        people.setdefault(person, role)
        for w in weeks:
            cell = ws.cell(row, FIRST_WEEK_COL + w['index'])
            val = cell.value
            phase = PHASE_BY_FILL.get(fill_key(cell))
            hours = None
            if isinstance(val, (int, float)):
                hours = float(val)
            elif isinstance(val, str) and 'vacac' in val.lower():
                phase = 'vacaciones'
            elif val is None and phase is None:
                continue
            if val is None and phase != 'vacaciones':
                continue  # celda sólo coloreada, sin dato
            rows.append({'client': client, 'person': person, 'week': w['index'],
                         'hours': hours, 'phase': phase})
        row += 1

    out = {
        'source': 'PRESUPUESTO DE HORAS 2026.xlsx',
        'period': {'from': weeks[0]['start'], 'to': weeks[-1]['start']},
        'weeks': weeks,
        'people': [{'name': n, 'role': r} for n, r in people.items()],
        'clients': clients,
        'budget': rows,
    }
    with open('data/presupuesto-2026.json', 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
    print('semanas', len(weeks), 'personas', len(people), 'clientes', len(clients), 'celdas', len(rows))


if __name__ == '__main__':
    main(sys.argv[1])
